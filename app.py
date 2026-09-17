from flask import Flask,jsonify,render_template,request
import random,math,datetime,os,io,re,csv
from zoneinfo import ZoneInfo
from urllib.parse import urlsplit,urlunsplit,parse_qsl,urlencode
import requests
from openpyxl import load_workbook

app=Flask(__name__); TZ=ZoneInfo('America/Cancun')
BASE={'seed':20260917,'admissions_day':30.5,'births_day':7.5,'cesarean_probability':.41,'admission_probability':.55,'floor_capacity':27,'initial_floor_occupancy':22,'labor_capacity':9,'initial_labor_occupancy':4,'physicians_per_shift':4,'triage_minutes':35}
PHEDS=[]
for i in range(1,25): PHEDS.append({'id':str(i),'area':'Hospitalización GO','oxygen':True,'isolation':False,'sex':'MUJER','censable':True,'fictitious':False})
for i in range(25,28): PHEDS.append({'id':str(i),'area':'Hospitalización GO · Aislamiento','oxygen':True,'isolation':True,'sex':'MUJER','censable':True,'fictitious':False})
for i in range(1,11): PHEDS.append({'id':f'T{i:02d}','area':'Tococirugía','oxygen':i!=10,'isolation':False,'sex':'MUJER','censable':False,'fictitious':False})
PHEDS += [{'id':'E01','area':'Expulsión','oxygen':True,'isolation':False,'sex':'MUJER','censable':False,'fictitious':False},{'id':'M01','area':'Sala Mixta','oxygen':True,'isolation':False,'sex':'MUJER','censable':False,'fictitious':False},{'id':'UT01','area':'Urgencias Tococirugía','oxygen':True,'isolation':False,'sex':'MUJER','censable':False,'fictitious':True}]
for i in range(1,6): PHEDS.append({'id':f'RP{i:02d}','area':'Recuperación Postparto','oxygen':i!=5,'isolation':False,'sex':'MUJER','censable':False,'fictitious':False})

# Plantilla pública pseudonimizada. Nunca se exponen cédulas ni identificadores personales.
STAFF=[
{'id':'M01','shift':'TED','schedule':'07:00–19:00 sábados, domingos y feriados','area':'Triage · Tococirugía · Piso','vacations':['28 mar–12 abr','9–24 ene 2026']},
{'id':'M02','shift':'TED','schedule':'07:00–19:00 sábados, domingos y feriados','area':'Triage · Tococirugía · Piso','vacations':['14 feb–1 mar','19 dic–10 ene 2027','5–19 sep']},
{'id':'M03','shift':'TED','schedule':'07:00–19:00 sábados, domingos y feriados','area':'Medicina materno fetal','vacations':['3–17 may','7–21 nov']},
{'id':'M04','shift':'TEN','schedule':'19:00–07:00 sábados, domingos y feriados','area':'Triage · Tococirugía','vacations':['28 mar–12 abr','19 dic–1 ene','25 jul–9 ago']},
{'id':'M05','shift':'TEN','schedule':'19:00–07:00 sábados, domingos y feriados','area':'Triage · Tococirugía','vacations':['3–17 may','28 nov–13 dic','4–19 jul']},
{'id':'M06','shift':'TM','schedule':'07:00–15:00 lunes a viernes','area':'Triage · Tococirugía','vacations':['1–14 abr','17–31 dic','14–30 sep']},
{'id':'M07','shift':'TM','schedule':'07:00–15:00 lunes a viernes','area':'Triage · Tococirugía','vacations':['15–28 abr','6–17 jul','30 nov–15 dic']},
{'id':'M08','shift':'TM','schedule':'07:00–15:00 lunes a viernes','area':'Triage · Tococirugía','vacations':['18–31 mar','5–16 oct']},
{'id':'M09','shift':'TM','schedule':'07:00–15:00 lunes a viernes','area':'Consulta externa · Piso · Colposcopia','vacations':['18–29 may','3–14 ago']},
{'id':'M10','shift':'TNA','schedule':'20:30–07:30 nocturno MJ-LMV','area':'Triage · Tococirugía · Piso','vacations':['6–16 abr','18–30 dic','21 jul–4 ago']},
{'id':'M11','shift':'TNA','schedule':'19:00–07:00 nocturno MJ-LMV','area':'Triage · Tococirugía · Piso','vacations':['18–31 mar','17–30 nov','6–17 jul']},
{'id':'M12','shift':'TNA','schedule':'21:00–09:00 MJ-LXV nocturno','area':'Triage · Tococirugía · Piso','vacations':['20–30 ene 2027','6–18 ago','18 nov–2 dic']},
{'id':'M13','shift':'TNB','schedule':'21:00–08:00 nocturno LMV-MJ','area':'Triage · Tococirugía · Piso','vacations':['17–31 dic','14–30 sep']},
{'id':'M14','shift':'TNB','schedule':'20:00–08:00 nocturno LMV-MJ','area':'Triage · Tococirugía · Piso','vacations':[]},
{'id':'M15','shift':'TNB','schedule':'20:00–07:00 nocturno LMV-MJ','area':'Triage · Tococirugía · Piso','vacations':['21,25,27,31 ago–2 sep','1–11 dic','3–13 nov']},
{'id':'M16','shift':'TV','schedule':'14:00–22:00 lunes a viernes','area':'Triage','vacations':['1–14 abr','17–31 dic','3–14 ago']},
{'id':'M17','shift':'TV','schedule':'13:30–21:00 lunes a viernes','area':'Triage · Tococirugía','vacations':['30 abr–15 may','17–30 nov']},
{'id':'M18','shift':'TE','schedule':'Sáb 07:00–19:00 · Dom–Lun 07:00–07:00','area':'Triage · Tococirugía · Piso','vacations':[]},
{'id':'M19','shift':'TM','schedule':'07:00–15:00 lunes a viernes','area':'Oncología ginecológica','vacations':[]}
]

def now(): return datetime.datetime.now(TZ)
def today(): return now().date()
def staff_status():
    d=today(); out=[]
    for x in STAFF:
        y=dict(x); y['status']='PROGRAMADO'
        if d.year==2026 and d.month==9:
            if x['id']=='M02' and 5<=d.day<=19:y['status']='VACACIONES'
            if x['id'] in ('M06','M13') and 14<=d.day<=30:y['status']='VACACIONES'
        out.append(y)
    return out

def normdate(v):
    if isinstance(v,datetime.datetime): return v.date()
    if isinstance(v,datetime.date): return v
    if not v:return None
    s=str(v).strip()
    for f in ('%d/%m/%Y','%d/%m/%y','%Y-%m-%d','%d-%m-%Y','%d-%m-%y'):
        try:return datetime.datetime.strptime(s,f).date()
        except:pass
    return None

def normtime(v):
    if isinstance(v,datetime.datetime): return v.time()
    if isinstance(v,datetime.time): return v
    if isinstance(v,(int,float)) and 0<=v<1:
        sec=round(v*86400)%86400; return datetime.time(sec//3600,(sec%3600)//60,sec%60)
    if not v:return None
    s=str(v).strip().replace(';',':').replace(' ', '')
    s=re.sub(r'\.(?=\d{2}$)',':',s)
    for f in ('%H:%M','%H:%M:%S'):
        try:return datetime.datetime.strptime(s,f).time()
        except:pass
    return None

def with_download(url):
    p=urlsplit(url.strip()); q=dict(parse_qsl(p.query,keep_blank_values=True)); q['download']='1'
    return urlunsplit((p.scheme,p.netloc,p.path,urlencode(q),p.fragment))

def getbytes(url):
    if not url:return None
    urls=[url]
    if 'onedrive.live.com' in url or '1drv.ms' in url or 'sharepoint.com' in url: urls.insert(0,with_download(url))
    last=''
    for u in urls:
        try:
            r=requests.get(u,timeout=25,allow_redirects=True,headers={'User-Agent':'Mozilla/5.0'}); r.raise_for_status()
            data=r.content; ctype=(r.headers.get('content-type') or '').lower()
            if data[:2]==b'PK': return data
            last=f'contenido no XLSX ({ctype or "sin content-type"})'
        except Exception as e:last=str(e)
    raise ValueError('OneDrive/SharePoint no entregó un XLSX descargable: '+last[:120])

def sheet_csv(sheet_id,gid):
    r=requests.get(f'https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}',timeout=15); r.raise_for_status(); return r.text

def live_births():
    try:
        sid=os.getenv('BIRTHS_SHEET_ID','1SJGGbbUGA2pTmouDi-tOFi_kVmtSVFhGFXBq70kFOTw'); gid=os.getenv('BIRTHS_GID','405341237')
        rows=list(csv.reader(io.StringIO(sheet_csv(sid,gid)))); vals=[]
        for row in rows[1:]:
            if len(row)>5:
                d=normdate(row[3]); typ=row[5].strip().lower()
                if d and d<=today() and typ in ('parto','cesarea','cesárea'): vals.append((d,typ))
        if not vals:return None
        last=max(d for d,_ in vals); day=[x for x in vals if x[0]==last]
        return {'value':len(day),'vaginal':sum(t=='parto' for _,t in day),'cesareans':sum(t!='parto' for _,t in day),'last_valid_date':last.isoformat(),'status':'REAL · EN VIVO'}
    except Exception as e:return {'value':None,'last_valid_date':None,'status':'SIN CONEXIÓN','error':str(e)[:160]}

def parse_census(wb):
    candidates=[]
    for ws in wb.worksheets:
        m=re.fullmatch(r'0?(\d{1,2})',ws.title.strip())
        if m:
            try:d=datetime.date(today().year,today().month,int(m.group(1)))
            except:continue
            if d<=today():candidates.append((d,ws))
    if not candidates:return None
    d,ws=max(candidates,key=lambda x:x[0])
    expected={str(300+i):f'{i:02d}' for i in range(1,25)}
    expected.update({f'AISL{i}':f'{24+i:02d}' for i in range(1,4)})
    seen={}; labor=0
    for row in ws.iter_rows(values_only=True):
        if len(row)<3:continue
        raw=row[0]
        if isinstance(raw,(int,float)) and float(raw).is_integer():bed=str(int(raw))
        else:bed=str(raw or '').strip().upper()
        service=str(row[1] or '').strip().upper()
        if bed in expected and bed not in seen and 'GINECO' in service:
            seen[bed]=bool(str(row[2] or '').strip())
        txt=' | '.join(str(x or '') for x in row[:8]).upper()
        if ('LABOR' in txt or 'TOCOCIR' in txt) and any(str(x or '').strip() for x in row[2:5]):labor+=1
    occupied=sorted((expected[k] for k,v in seen.items() if v),key=int)
    return {'value':len(occupied),'occupied_beds':occupied,'labor':labor,'last_valid_date':d.isoformat(),'status':'REAL · EN VIVO'}

def find_header(ws):
    for rn,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,40),values_only=True),1):
        labels=[str(x or '').strip().upper() for x in row]
        if any('FOLIO' in x for x in labels) and any(x=='FECHA' or x.startswith('FECHA ') for x in labels) and any('HORA' in x for x in labels):
            return rn,labels
    return None,None

def parse_triage(wb):
    events=[]
    for ws in wb.worksheets:
        hr,labels=find_header(ws)
        if not hr:continue
        fi=next((i for i,x in enumerate(labels) if 'FOLIO' in x),None)
        di=next((i for i,x in enumerate(labels) if x=='FECHA' or x.startswith('FECHA ')),None)
        ti=next((i for i,x in enumerate(labels) if 'HORA' in x),None)
        if di is None:continue
        for row in ws.iter_rows(min_row=hr+1,values_only=True):
            d=normdate(row[di] if di<len(row) else None)
            if not d or d>today():continue
            fol=str(row[fi] or '').strip() if fi is not None and fi<len(row) else ''
            tm=normtime(row[ti] if ti is not None and ti<len(row) else None)
            events.append((d,fol,tm))
    if not events:return None
    last=max(x[0] for x in events); day=[x for x in events if x[0]==last]
    dedup={}
    for i,x in enumerate(day):dedup[(x[0],x[1] or f'ROW{i}')]=x
    day=list(dedup.values()); shifts={'Matutino':0,'Vespertino':0,'Nocturno':0}; by_hour={}
    for _,_,tm in day:
        if not tm:continue
        h=tm.hour; by_hour[str(h).zfill(2)]=by_hour.get(str(h).zfill(2),0)+1
        mins=h*60+tm.minute
        if 420<=mins<=840:shifts['Matutino']+=1
        elif 841<=mins<=1260:shifts['Vespertino']+=1
        else:shifts['Nocturno']+=1
    return {'value':len(day),'by_shift':shifts,'by_hour':by_hour,'last_valid_date':last.isoformat(),'status':'REAL · EN VIVO'}

def live_xlsx(url,kind):
    try:
        data=getbytes(url)
        if not data:return None
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        result=parse_census(wb) if kind=='census' else parse_triage(wb)
        return result or {'value':None,'last_valid_date':None,'status':'SIN DATOS VÁLIDOS'}
    except Exception as e:return {'value':None,'last_valid_date':None,'status':'SIN CONEXIÓN','error':str(e)[:160]}

def real_state():
    b=live_births(); c=live_xlsx(os.getenv('CENSUS_XLSX_URL'),'census'); t=live_xlsx(os.getenv('TRIAGE_XLSX_URL'),'triage')
    src=[{'name':'Nacimientos 2026','status':b['status'] if b else 'SIN CONEXIÓN','last_valid_date':b.get('last_valid_date') if b else None},{'name':'Censo Piso / Labor','status':c['status'] if c else 'CONFIGURAR URL','last_valid_date':c.get('last_valid_date') if c else None},{'name':'Urgencias / Triage','status':t['status'] if t else 'CONFIGURAR URL','last_valid_date':t.get('last_valid_date') if t else None},{'name':'Infraestructura PHEDS','status':'REAL · MAESTRO','last_valid_date':today().isoformat()},{'name':'Plantilla médica','status':'REAL · MAESTRO','last_valid_date':today().isoformat()}]
    return {'mode':'REAL','as_of':now().isoformat(),'cutoff_date':today().isoformat(),'sources':src,'births':b or {'value':None},'floor':{'value':c.get('value') if c else None,'capacity':27,'occupied_beds':c.get('occupied_beds',[]) if c else [],'last_valid_date':c.get('last_valid_date') if c else None},'labor':{'value':c.get('labor') if c else None,'capacity':9,'last_valid_date':c.get('last_valid_date') if c else None},'triage':{'value':t.get('value') if t else None,'by_shift':t.get('by_shift',{}) if t else {},'by_hour':t.get('by_hour',{}) if t else {},'last_valid_date':t.get('last_valid_date') if t else None},'infrastructure':{'total':len(PHEDS),'censable':sum(x['censable'] for x in PHEDS),'functional':sum(not x['censable'] for x in PHEDS)}}

def poisson(lam,rng):
    l=math.exp(-lam);k=0;q=1
    while q>l:k+=1;q*=rng.random()
    return k-1

def simulate(days=7,params=None):
    p=dict(BASE);p.update(params or {});rng=random.Random(int(p['seed']));rows=[];floor=p['initial_floor_occupancy'];labor=p['initial_labor_occupancy']
    for d in range(days):
        adm=poisson(float(p['admissions_day']),rng);births=poisson(float(p['births_day']),rng);cs=sum(rng.random()<p['cesarean_probability'] for _ in range(births));admitted=sum(rng.random()<p['admission_probability'] for _ in range(adm));labor=max(0,min(p['labor_capacity'],labor+min(admitted,births)-births));floor=max(0,min(p['floor_capacity'],floor+admitted-poisson(max(1,floor/1.2),rng)));rows.append({'day':d+1,'emergency_arrivals':adm,'hospital_admissions':admitted,'labor_occupancy':labor,'labor_capacity':p['labor_capacity'],'births':births,'vaginal':births-cs,'cesareans':cs,'floor_occupancy':floor,'floor_pct':round(100*floor/p['floor_capacity'],1),'triage_queue':max(0,round(adm/24*p['triage_minutes']/60-p['physicians_per_shift']))})
    return {'mode':'SIMULADO','days':rows,'summary':{'emergency_arrivals':sum(x['emergency_arrivals'] for x in rows),'births':sum(x['births'] for x in rows),'peak_floor_pct':max(x['floor_pct'] for x in rows),'peak_labor_occupancy':max(x['labor_occupancy'] for x in rows)},'parameters':p}

@app.route('/')
def home():return render_template('index.html')
@app.route('/health')
def health():return jsonify({'status':'ok','server_time':now().isoformat(),'pheds_positions':len(PHEDS),'physicians':len(STAFF)})
@app.route('/api/config')
def config():return jsonify(BASE)
@app.route('/api/pheds')
def pheds():return jsonify({'total':len(PHEDS),'censable':sum(x['censable'] for x in PHEDS),'functional':sum(not x['censable'] for x in PHEDS),'positions':PHEDS})
@app.route('/api/staff')
def staff():
    s=staff_status(); return jsonify({'total':len(s),'vacation':sum(x['status']=='VACACIONES' for x in s),'available_template':sum(x['status']!='VACACIONES' for x in s),'date':today().isoformat(),'physicians':s})
@app.route('/api/sources')
def sources():return jsonify(real_state()['sources'])
@app.route('/api/real-state')
def state():return jsonify(real_state())
@app.route('/api/simulate',methods=['POST'])
def sim():
    b=request.get_json(silent=True) or {};return jsonify(simulate(int(b.get('days',7)),b.get('parameters')))
@app.route('/api/predict')
def predict():return jsonify({'mode':'PREDICHO','plus_2h':{'floor_pct':0,'labor_pct':0,'triage_queue':0},'plus_4h':{'floor_pct':0,'labor_pct':0,'triage_queue':0},'plus_6h':{'floor_pct':0,'labor_pct':0,'triage_queue':0},'warning':'Predicción suspendida hasta disponer de las tres fuentes reales simultáneamente.'})
if __name__=='__main__':app.run(host='0.0.0.0',port=10000)
