import source_patch as patched
import ingestion_api
from app import normdate,normtime,today
from flask import jsonify,request
from openpyxl import load_workbook
import os,io,re,datetime,csv,requests,unicodedata
app=patched.app; getbytes=patched.getbytes
ingestion_api.register(app, getbytes)
EXPECTED={str(300+i):f'{i:02d}' for i in range(1,25)};EXPECTED.update({f'AISL{i}':f'{24+i:02d}' for i in range(1,4)})
def clean(s):return ''.join(c for c in unicodedata.normalize('NFD',str(s or '').upper()) if unicodedata.category(c)!='Mn')
def bedkey(raw):
 if isinstance(raw,(int,float)) and float(raw).is_integer():return str(int(raw))
 return re.sub(r'\s+','',clean(raw))
def hdp(dx):
 s=clean(dx);compact=re.sub(r'[^A-Z]','',s)
 # Solo diagnóstico registrado. Se toleran errores ortográficos frecuentes del censo.
 if 'HELLP' in compact:return 'hellp'
 if re.search(r'\bECLAMPS',s) and not re.search(r'PRE+\w*E?CLAMPS',s):return 'eclampsia'
 if any(x in compact for x in ('PREECLAMPS','PREEVCLAMPS','PRECLAMPS','PREECLAPS','PRECLAPS')):return 'preeclampsia'
 if (('HIPERTENSION' in s or re.search(r'\bHTA\b',s)) and ('GESTACIONAL' in s or 'GESTACION' in s or 'EMBARAZ' in s)):return 'hipertension_gestacional'
 if re.search(r'\bE\s*\.?\s*H\s*\.?\s*E\s*\.?\b',s):return 'the_no_especificado'
 return None
def go_block_rows(ws):
 # Localiza el bloque físico de hospitalización G/O por camas 301-324; evita AISL de otros servicios.
 rows=list(ws.iter_rows(values_only=True));start=None
 for i,row in enumerate(rows):
  if bedkey(row[0] if row else None)=='301' and len(row)>1 and 'GINECO' in clean(row[1]):start=i;break
 if start is None:return []
 out=[]
 for row in rows[start:]:
  k=bedkey(row[0] if row else None)
  if k in EXPECTED:out.append(row)
  if k=='AISL3' and len(out)>=20:break
 return out
def census_sheet(ws,d):
 seen={};cats={k:[] for k in ('preeclampsia','eclampsia','hipertension_gestacional','hellp','the_no_especificado')}
 for row in go_block_rows(ws):
  if len(row)<3:continue
  k=bedkey(row[0]);service=clean(row[1])
  if k in EXPECTED and k not in seen and 'GINECO' in service:
   occupied=bool(str(row[2] or '').strip());seen[k]=occupied
   if occupied:
    # En este formato DIAGNOSTICOS está en columna G (índice 6).
    diagnosis=row[6] if len(row)>6 else None;cat=hdp(diagnosis)
    if cat:cats[cat].append(EXPECTED[k])
 if len(seen)<20:return None
 occ=sorted((EXPECTED[k] for k,v in seen.items() if v),key=int)
 for k in cats:cats[k]=sorted(set(cats[k]),key=int)
 allbeds=sorted(set(sum(cats.values(),[])),key=int)
 hypertension={'total':len(allbeds),'beds':allbeds,'preeclampsia':len(cats['preeclampsia']),'preeclampsia_beds':cats['preeclampsia'],'eclampsia':len(cats['eclampsia']),'eclampsia_beds':cats['eclampsia'],'hipertension_gestacional':len(cats['hipertension_gestacional']),'hipertension_gestacional_beds':cats['hipertension_gestacional'],'hellp':len(cats['hellp']),'hellp_beds':cats['hellp'],'the_no_especificado':len(cats['the_no_especificado']),'the_no_especificado_beds':cats['the_no_especificado'],'method':'DIAGNÓSTICO REGISTRADO EN CENSO · LECTOR G/O 301-324 + AISL1-3'}
 return {'date':d.isoformat(),'occupied':len(occ),'available':27-len(occ),'occupancy_pct':round(100*len(occ)/27,1),'occupied_beds':occ,'hypertensive_disorders':hypertension,'status':'REAL'}
def census_history(days):
 data=getbytes(os.getenv('CENSUS_XLSX_URL'));wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True);out={};start=today()-datetime.timedelta(days=days-1)
 for ws in wb.worksheets:
  m=re.fullmatch(r'0?(\d{1,2})',ws.title.strip())
  if not m:continue
  try:d=datetime.date(today().year,today().month,int(m.group(1)))
  except:continue
  if start<=d<=today():
   x=census_sheet(ws,d)
   if x:out[d.isoformat()]=x
 return out
def triage_history(days):
 data=getbytes(os.getenv('TRIAGE_XLSX_URL'));wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True);events={};start=today()-datetime.timedelta(days=days-1)
 for ws in wb.worksheets:
  header=None
  for rn,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,40),values_only=True),1):
   labels=[str(x or '').strip().upper() for x in row]
   if any('FOLIO' in x for x in labels) and any(x=='FECHA' or x.startswith('FECHA ') for x in labels) and any('HORA' in x for x in labels):header=(rn,labels);break
  if not header:continue
  hr,labels=header;fi=next((i for i,x in enumerate(labels) if 'FOLIO' in x),None);di=next((i for i,x in enumerate(labels) if x=='FECHA' or x.startswith('FECHA ')),None);ti=next((i for i,x in enumerate(labels) if 'HORA' in x),None)
  for n,row in enumerate(ws.iter_rows(min_row=hr+1,values_only=True)):
   d=normdate(row[di] if di is not None and di<len(row) else None)
   if not d or not(start<=d<=today()):continue
   fol=str(row[fi] or '').strip() if fi is not None and fi<len(row) else f'R{n}';tm=normtime(row[ti] if ti is not None and ti<len(row) else None);events.setdefault(d.isoformat(),{})[fol or f'R{n}']=tm
 out={}
 for ds,items in events.items():
  shifts={'Matutino':0,'Vespertino':0,'Nocturno':0}
  for tm in items.values():
   if tm:
    mins=tm.hour*60+tm.minute;shifts['Matutino' if 420<=mins<=840 else 'Vespertino' if 841<=mins<=1260 else 'Nocturno']+=1
  out[ds]={'value':len(items),'by_shift':shifts,'status':'REAL'}
 return out
def births_history(days):
 sid=os.getenv('BIRTHS_SHEET_ID','1SJGGbbUGA2pTmouDi-tOFi_kVmtSVFhGFXBq70kFOTw');gid=os.getenv('BIRTHS_GID','405341237');r=requests.get(f'https://docs.google.com/spreadsheets/d/{sid}/export?format=csv&gid={gid}',timeout=20);r.raise_for_status();start=today()-datetime.timedelta(days=days-1);out={}
 for row in list(csv.reader(io.StringIO(r.text)))[1:]:
  if len(row)<=5:continue
  d=normdate(row[3]);typ=row[5].strip().lower()
  if not d or not(start<=d<=today()) or typ not in ('parto','cesarea','cesárea'):continue
  x=out.setdefault(d.isoformat(),{'value':0,'vaginal':0,'cesareans':0,'status':'REAL'});x['value']+=1;x['vaginal' if typ=='parto' else 'cesareans']+=1
 return out
@app.route('/api/history')
def history():
 try:days=max(1,min(int(request.args.get('days',7)),31))
 except:days=7
 dates=[(today()-datetime.timedelta(days=i)).isoformat() for i in range(days)];errors={}
 try:c=census_history(days)
 except Exception as e:c={};errors['census']=str(e)[:140]
 try:t=triage_history(days)
 except Exception as e:t={};errors['triage']=str(e)[:140]
 try:b=births_history(days)
 except Exception as e:b={};errors['births']=str(e)[:140]
 rows=[{'date':ds,'floor':c.get(ds,{'status':'SIN DATO'}),'triage':t.get(ds,{'status':'SIN DATO'}),'births':b.get(ds,{'status':'SIN DATO'})} for ds in dates]
 return jsonify({'mode':'REAL','days':rows,'errors':errors,'privacy':'SIN IDENTIFICADORES DIRECTOS'})
if __name__=='__main__':app.run(host='0.0.0.0',port=10000)
